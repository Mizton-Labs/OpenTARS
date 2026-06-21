"""DOCX extraction using python-docx."""

from __future__ import annotations

import io
import logging

logger = logging.getLogger(__name__)

_DOCX_VERSION: str | None = None

try:
    import docx  # type: ignore[import-untyped]  # python-docx

    _DOCX_VERSION = docx.__version__
except ImportError:  # pragma: no cover
    docx = None  # type: ignore[assignment]


def is_available() -> bool:
    return docx is not None


def extract_docx(data: bytes) -> tuple[str, str, list[str]]:
    """Extract text from DOCX bytes using python-docx.

    Returns (extracted_text, parser_version, warnings).
    """
    if docx is None:
        raise RuntimeError("python-docx is not installed")

    warnings: list[str] = []
    version = _DOCX_VERSION or "unknown"

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"Cannot open DOCX: {exc}") from exc

    paragraphs: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # Also extract table content
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cells)
            if line.strip("| "):
                paragraphs.append(line)

    if not paragraphs:
        warnings.append("No text content found in DOCX")

    return "\n\n".join(paragraphs), version, warnings
